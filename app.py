"""نظام حجز مواعيد تقديم طلبات الدمج — تطبيق Flask خفيف بقاعدة SQLite."""
from __future__ import annotations

import hmac
import io
import os
import re
import secrets
import sqlite3
import time as _time
from datetime import date, datetime, timedelta, timezone
from functools import wraps
from pathlib import Path

from flask import (
    Flask,
    abort,
    flash,
    g,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from barcode import Code128
from barcode.writer import SVGWriter
from markupsafe import Markup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.middleware.proxy_fix import ProxyFix

BASE_DIR = Path(__file__).resolve().parent
# على Railway تُحفظ البيانات في الـ Volume المرفق تلقائيًا
DATA_DIR = Path(
    os.environ.get("DATA_DIR")
    or os.environ.get("RAILWAY_VOLUME_MOUNT_PATH")
    or BASE_DIR / "data"
)
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = DATA_DIR / "booking.db"


def _load_env_file() -> None:
    env_file = BASE_DIR / ".env"
    if not env_file.exists():
        return
    for line in env_file.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def _secret_key() -> str:
    if os.environ.get("SECRET_KEY"):
        return os.environ["SECRET_KEY"]
    key_file = DATA_DIR / "secret.key"
    if not key_file.exists():
        key_file.write_text(secrets.token_hex(32), encoding="utf-8")
    return key_file.read_text(encoding="utf-8").strip()


_load_env_file()

ON_RAILWAY = bool(os.environ.get("RAILWAY_ENVIRONMENT"))

app = Flask(__name__)
# خلف بروكسي الاستضافة: خذ البروتوكول (https) والعنوان من ترويسات البروكسي
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
app.config.update(
    SECRET_KEY=_secret_key(),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE", "1" if ON_RAILWAY else "0") == "1",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=12),
    MAX_CONTENT_LENGTH=64 * 1024,
)

# ============================================================
#  البيانات المرجعية (مطابقة للنظام الأساسي)
# ============================================================
ADMINISTRATIONS = [
    "بنها", "غرب شبرا الخيمة", "شرق شبرا الخيمة", "قها", "طوخ", "كفر شكر",
    "قليوب", "شبين القناطر", "العبور", "الخانكة", "الخصوص", "القناطر الخيرية",
]

ACADEMIC_YEAR_LABEL = "2026م / 2027م"

# شروط قبول الدمج: الاسم -> (أدنى ذكاء، أقصى ذكاء، الشروط، الأوراق الإضافية الخاصة بالحالة)
DISABILITY_RULES = {
    "إعاقة ذهنية بسيطة": (
        None, None,
        "تقرير طبي شامل للحالة + مقياس ستانفورد بينيه للذكاء.", []),
    "بطء التعلم": (
        71, 84,
        "تقرير طبي شامل للحالة + مقياس ستانفورد بينيه للذكاء (71 – 84).", []),
    "متلازمة داون": (
        60, None,
        "تقرير طبي شامل موضح به احتياج الطالب إلى مرافق من عدمه + "
        "مقياس ستانفورد بينيه للذكاء 60 فأكثر.", []),
    "الشلل الدماغي (CP)": (
        60, None,
        "تقرير طبي شامل موضح به احتياج الطالب إلى مرافق من عدمه + "
        "مقياس ستانفورد بينيه للذكاء 60 فأكثر.", []),
    "طيف التوحد": (
        None, None,
        "تقرير طبي شامل موضح به احتياج الطالب لمرافق من عدمه + مقياس ستانفورد بينيه "
        "للذكاء على الجانب الأدائي فقط + مقياس كارز أو مقياس جيليام.",
        ["مقياس كارز (CARS) أو مقياس جيليام (GARS)"]),
    "الإعاقة البصرية": (
        80, None,
        "قياس حدة إبصار بالكمبيوتر + تقرير طبي شامل مع توضيح نسبة الإبصار بالعينين "
        "طبقًا لمقياس الكمبيوتر + مقياس ستانفورد بينيه بدرجة 80 فأكثر.",
        ["قياس حدة الإبصار بالكمبيوتر"]),
    "الإعاقة السمعية": (
        80, None,
        "مقياس سمع بالسماعة بدرجة من (40 – 70) ديسيبل + تقرير طبي طبقًا لمقياس السمع + "
        "مقياس ستانفورد بينيه بدرجة 80 فأكثر. وفي حالات زارعي القوقعة يشترط تقرير طبي "
        "أن الطالب من حالات زارعي القوقعة ودرجة ذكاء 80 فأكثر.",
        ["مقياس السمع بالسماعة (40 – 70 ديسيبل)"]),
    "الإعاقة الحركية": (
        85, None,
        "تقرير طبي شامل للحالة + مقياس ستانفورد بينيه للذكاء بدرجة 85 فأكثر.", []),
    "صعوبات التعلم": (
        90, None,
        "اختبار الينوي أو مايكل بست أو مقياس أ.د عادل عبد الله أو بطارية فتحي الزيات + "
        "تقرير طبي شامل يتضمن مقياس ستانفورد بينيه بدرجة ذكاء 90 فأكثر.",
        ["اختبار صعوبات التعلم (الينوي / مايكل بست / عادل عبد الله / فتحي الزيات)"]),
    "فرط الحركة وتشتت الانتباه": (
        80, 85,
        "اختبار كونرز بدرجة 65 فأكثر + تقرير طبي شامل للحالة متضمن اختبار ذكاء "
        "ستانفورد بينيه بدرجة من (80 – 85).",
        ["اختبار كونرز (65 فأكثر)"]),
    "غير ذلك": (None, None, "", []),
}

IMPORTANT_NOTES = [
    "في حالة اقتران الإعاقة الحركية بأي إعاقة أخرى فإنها لا تُعد إعاقة حيث إنها لا تؤثر على "
    "التحصيل الدراسي، ويُصنف الطالب وفقًا لدرجة الذكاء من (60 – 84).",
    "يتم مراعاة التوصيات الواردة في التقارير الطبية بشأن الفصل الدراسي، لجنة بالدور الأرضي / "
    "مرافق للكتابة ... حسب درجة الإعاقة الحركية.",
]


def iq_range_label(disability: str) -> str | None:
    low, high = DISABILITY_RULES[disability][:2]
    if low is not None and high is not None:
        return f"من {low} إلى {high}"
    if low is not None:
        return f"{low} فأكثر"
    return None


def check_iq(disability: str, raw: str) -> tuple[int | None, str | None]:
    """يعيد (الدرجة، رسالة الخطأ) طبقًا لشروط قبول نوع الإعاقة."""
    low, high = DISABILITY_RULES[disability][:2]
    label = iq_range_label(disability)
    if not raw:
        if disability == "غير ذلك":
            return None, None
        return None, "درجة الذكاء على مقياس ستانفورد بينيه مطلوبة" + (
            f" لحالات «{disability}» ({label})." if label else "."
        )
    iq = int(raw)
    if iq > 200:
        return None, "درجة الذكاء يجب أن تكون بين 0 و200."
    if (low is not None and iq < low) or (high is not None and iq > high):
        return None, (
            f"درجة الذكاء ({iq}) لا تطابق شروط قبول «{disability}» للعام {ACADEMIC_YEAR_LABEL}: "
            f"المطلوب درجة ذكاء {label} على مقياس ستانفورد بينيه."
        )
    return iq, None

STAGES = [
    ("رياض الأطفال", ["KG1", "KG2"]),
    ("ابتدائي", [
        "الأول الابتدائي", "الثاني الابتدائي", "الثالث الابتدائي",
        "الرابع الابتدائي", "الخامس الابتدائي", "السادس الابتدائي",
    ]),
    ("إعدادي", ["الأول الإعدادي", "الثاني الإعدادي", "الثالث الإعدادي"]),
    ("ثانوي عام", ["الأول الثانوي العام", "الثاني الثانوي العام", "الثالث الثانوي العام"]),
]
GRADE_STAGE = {grade: stage for stage, grades in STAGES for grade in grades}

RELATIONS = [
    "الأب", "الأم", "الأخ", "الأخت", "الجد", "الجدة",
    "العم أو الخال", "العمة أو الخالة", "وصي قانوني", "أخرى",
]

GENERAL_DOCUMENTS = [
    "صورتان شخصيتان للطالب",
    "صورة شهادة ميلاد الطالب",
    "صورة بطاقة الرقم القومي للأب",
    "إثبات القيد من المدرسة",
    "خطاب الأخصائي النفسي من المدرسة",
    "إقرار ولي الأمر",
    "تقرير طبي شامل للحالة",
    "مقياس ستانفورد بينيه للذكاء",
]

# Python weekday(): الاثنين=0 ... الأحد=6
WEEKDAY_NAMES = {
    5: "السبت", 6: "الأحد", 0: "الاثنين", 1: "الثلاثاء",
    2: "الأربعاء", 3: "الخميس", 4: "الجمعة",
}
WEEKDAY_ORDER = [5, 6, 0, 1, 2, 3, 4]
MONTH_NAMES = [
    "يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
    "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر",
]

DEFAULT_SETTINGS = {
    "weekdays": "6,1,3",  # أحد، ثلاثاء، خميس
    "capacity": "200",
    "start": "09:00",
    "end": "14:00",
    "slot_minutes": "15",
    "is_open": "1",
    "closed_message": "",
    "open_at": "",
    "announcement": "",
}
LOOKAHEAD_DAYS = 120

SCHEMA = """
CREATE TABLE IF NOT EXISTS settings (
    key TEXT PRIMARY KEY,
    value TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS day_overrides (
    day TEXT PRIMARY KEY,
    capacity INTEGER,
    closed INTEGER NOT NULL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS registrations (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    token TEXT NOT NULL UNIQUE,
    code TEXT NOT NULL UNIQUE,
    day TEXT NOT NULL,
    seq INTEGER NOT NULL,
    slot_start TEXT NOT NULL,
    slot_end TEXT NOT NULL,
    student_name TEXT NOT NULL,
    student_nid TEXT NOT NULL,
    birth_date TEXT NOT NULL,
    gender TEXT NOT NULL,
    administration TEXT NOT NULL,
    school TEXT NOT NULL,
    stage TEXT NOT NULL,
    grade TEXT NOT NULL,
    disability TEXT NOT NULL,
    guardian_name TEXT NOT NULL,
    relation TEXT NOT NULL,
    phone TEXT NOT NULL,
    whatsapp TEXT,
    guardian_nid TEXT,
    created_at TEXT NOT NULL,
    UNIQUE (day, seq)
);
CREATE INDEX IF NOT EXISTS ix_registrations_day ON registrations (day);
CREATE INDEX IF NOT EXISTS ix_registrations_nid ON registrations (student_nid);
"""

# أعمدة أُضيفت بعد الإصدار الأول — تُضاف تلقائيًا لقواعد البيانات القديمة
EXTRA_COLUMNS = {
    "iq_score": "INTEGER",
    "condition_description": "TEXT",
    "report_date": "TEXT",
    "report_authority": "TEXT",
    "city": "TEXT",
    "address": "TEXT",
    "email": "TEXT",
    "father_name": "TEXT",
    "father_nid": "TEXT",
    "mother_name": "TEXT",
    "mother_nid": "TEXT",
    "status": "TEXT NOT NULL DEFAULT 'active'",
    "stopped_at": "TEXT",
    "received_at": "TEXT",
}

STATUS_ACTIVE, STATUS_STOPPED = "active", "stopped"


# ============================================================
#  قاعدة البيانات
# ============================================================
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
# على Railway: قاعدة Postgres دائمة لا تُمسح مع إعادة النشر. بدونها: ملف SQLite محلي.
USE_POSTGRES = DATABASE_URL.startswith(("postgres://", "postgresql://"))
PG_SCHEMA = os.environ.get("PG_SCHEMA", "booking")
PG_LOCK_KEY = 470_2026  # قفل تسلسل الحجز (بديل BEGIN IMMEDIATE في SQLite)

INTEGRITY_ERRORS: tuple = (sqlite3.IntegrityError,)
if USE_POSTGRES:
    import psycopg

    INTEGRITY_ERRORS += (psycopg.IntegrityError,)


class _Row(dict):
    """صف يدعم الوصول بالاسم وبالرقم مثل sqlite3.Row."""

    def __init__(self, columns, values):
        super().__init__(zip(columns, values))
        self._values = tuple(values)

    def __getitem__(self, key):
        if isinstance(key, int):
            return self._values[key]
        return super().__getitem__(key)

    def __iter__(self):
        return iter(self._values)


def _pg_row_factory(cursor):
    columns = [c.name for c in cursor.description] if cursor.description else []
    return lambda values: _Row(columns, values)


class PgConnection:
    """واجهة مطابقة لما يستخدمه التطبيق من sqlite3.Connection."""

    def __init__(self):
        self._conn = psycopg.connect(DATABASE_URL, autocommit=True, connect_timeout=10)
        self._conn.execute(f"SET search_path TO {PG_SCHEMA}")

    def execute(self, sql: str, params=()):
        if sql.strip().upper() == "BEGIN IMMEDIATE":
            self._conn.execute("BEGIN")
            return self._conn.execute("SELECT pg_advisory_xact_lock(%s)", (PG_LOCK_KEY,))
        cursor = self._conn.cursor(row_factory=_pg_row_factory)
        cursor.execute(sql.replace("?", "%s"), params or None)
        return cursor

    def executescript(self, script: str) -> None:
        self._conn.execute(script)

    @property
    def in_transaction(self) -> bool:
        return self._conn.info.transaction_status != psycopg.pq.TransactionStatus.IDLE

    def close(self) -> None:
        self._conn.close()


def connect():
    if USE_POSTGRES:
        return PgConnection()
    conn = sqlite3.connect(DB_PATH, timeout=15, isolation_level=None)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    return conn


def get_db():
    if "db" not in g:
        g.db = connect()
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db_postgres() -> None:
    conn = connect()
    conn.executescript(
        f"CREATE SCHEMA IF NOT EXISTS {PG_SCHEMA};"
        + SCHEMA.replace("INTEGER PRIMARY KEY AUTOINCREMENT", "BIGSERIAL PRIMARY KEY")
    )
    for column, kind in EXTRA_COLUMNS.items():
        conn.execute(f"ALTER TABLE registrations ADD COLUMN IF NOT EXISTS {column} {kind}")
    for key, value in DEFAULT_SETTINGS.items():
        conn.execute(
            "INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT (key) DO NOTHING", (key, value)
        )
    conn.close()


def init_db() -> None:
    if USE_POSTGRES:
        init_db_postgres()
        return
    conn = connect()
    # قواعد بيانات الإصدار الأول كانت تمنع تكرار الرقم القومي نهائيًا؛ الآن يُسمح
    # بطلب جديد إذا كان السابق موقوفًا أو قديمًا، فيُعاد بناء الجدول بدون هذا القيد.
    old = conn.execute(
        "SELECT sql FROM sqlite_master WHERE type = 'table' AND name = 'registrations'"
    ).fetchone()
    legacy = old is not None and "student_nid TEXT NOT NULL UNIQUE" in old["sql"]
    if legacy:
        conn.execute("DROP INDEX IF EXISTS ix_registrations_day")
        conn.execute("ALTER TABLE registrations RENAME TO registrations_old")
    conn.executescript(SCHEMA)
    existing = {r["name"] for r in conn.execute("PRAGMA table_info(registrations)")}
    for column, kind in EXTRA_COLUMNS.items():
        if column not in existing:
            conn.execute(f"ALTER TABLE registrations ADD COLUMN {column} {kind}")
    if legacy:
        columns = ", ".join(r["name"] for r in conn.execute("PRAGMA table_info(registrations_old)"))
        conn.execute("BEGIN IMMEDIATE")
        conn.execute(
            f"INSERT INTO registrations ({columns}) SELECT {columns} FROM registrations_old"
        )
        conn.execute("DROP TABLE registrations_old")
        conn.execute("COMMIT")
    for key, value in DEFAULT_SETTINGS.items():
        conn.execute(
            "INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT (key) DO NOTHING", (key, value)
        )
    conn.close()


DEFAULT_CLOSED_MESSAGE = "التسجيل مغلق مؤقتًا، يرجى المحاولة في وقت آخر."


def save_settings(values: dict) -> None:
    conn = get_db()
    for key, value in values.items():
        conn.execute(
            "INSERT INTO settings (key, value) VALUES (?, ?)"
            " ON CONFLICT (key) DO UPDATE SET value = excluded.value",
            (key, value),
        )


def get_settings() -> dict:
    raw = dict(DEFAULT_SETTINGS)
    raw.update({r["key"]: r["value"] for r in get_db().execute("SELECT key, value FROM settings")})
    # الفتح التلقائي في الموعد الذي حدده المسؤول
    if raw["is_open"] != "1" and raw["open_at"] and raw["open_at"] <= now_cairo().strftime("%Y-%m-%dT%H:%M"):
        raw.update(is_open="1", open_at="")
        save_settings({"is_open": "1", "open_at": ""})
    return {
        "closed_message": raw["closed_message"],
        "open_at": raw["open_at"],
        "announcement": raw["announcement"],
        "weekdays": {int(x) for x in raw["weekdays"].split(",") if x.strip() != ""},
        "capacity": int(raw["capacity"]),
        "start": raw["start"],
        "end": raw["end"],
        "slot_minutes": int(raw["slot_minutes"]),
        "is_open": raw["is_open"] == "1",
    }


# ============================================================
#  الوقت والمواعيد
# ============================================================
def today_cairo() -> date:
    try:
        from zoneinfo import ZoneInfo

        return datetime.now(ZoneInfo("Africa/Cairo")).date()
    except Exception:
        return (datetime.now(timezone.utc) + timedelta(hours=2)).date()


def now_cairo() -> datetime:
    try:
        from zoneinfo import ZoneInfo

        return datetime.now(ZoneInfo("Africa/Cairo")).replace(tzinfo=None)
    except Exception:
        return datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(hours=2)


def now_cairo_str() -> str:
    return now_cairo().strftime("%Y-%m-%d %H:%M")


def _minutes(hhmm: str) -> int:
    hours, minutes = hhmm.split(":")
    return int(hours) * 60 + int(minutes)


def _hhmm(total: int) -> str:
    return f"{total // 60:02d}:{total % 60:02d}"


def slot_for(seq: int, capacity: int, settings: dict) -> tuple[str, str]:
    """توزيع أرقام اليوم بالتساوي على فترات العمل."""
    start, end = _minutes(settings["start"]), _minutes(settings["end"])
    step = settings["slot_minutes"]
    slots = max(1, (end - start) // step)
    capacity = max(capacity, 1)
    # الأرقام الزائدة عن السعة (بدل طلبات موقوفة) تُوزَّع من أول اليوم مرة أخرى
    index = min(slots - 1, ((seq - 1) % capacity) * slots // capacity)
    slot_start = start + index * step
    return _hhmm(slot_start), _hhmm(min(slot_start + step, end))


def day_capacity(conn, day: date, settings: dict) -> int:
    row = conn.execute(
        "SELECT capacity, closed FROM day_overrides WHERE day = ?", (day.isoformat(),)
    ).fetchone()
    if row is None:
        return settings["capacity"]
    if row["closed"]:
        return 0
    return row["capacity"] if row["capacity"] is not None else settings["capacity"]


def find_open_day(conn, settings: dict):
    """أول يوم استقبال قادم ما زال به أماكن: (اليوم، السعة) أو None."""
    if not settings["is_open"] or not settings["weekdays"]:
        return None
    today = today_cairo()
    for offset in range(1, LOOKAHEAD_DAYS + 1):
        day = today + timedelta(days=offset)
        if day.weekday() not in settings["weekdays"]:
            continue
        capacity = day_capacity(conn, day, settings)
        if capacity <= 0:
            continue
        taken = conn.execute(
            "SELECT COUNT(*) FROM registrations WHERE day = ? AND status = 'active'",
            (day.isoformat(),),
        ).fetchone()[0]
        if taken < capacity:
            return day, capacity
    return None


# ============================================================
#  التنسيق
# ============================================================
def fmt_day(value) -> str:
    day = date.fromisoformat(value) if isinstance(value, str) else value
    return f"{WEEKDAY_NAMES[day.weekday()]} {day.day} {MONTH_NAMES[day.month - 1]} {day.year}"


def fmt_datetime(value: str) -> str:
    """'2026-10-23T09:00' -> 'الجمعة 23 أكتوبر 2026 — 9:00 ص'"""
    day, _, clock = value.partition("T")
    return f"{fmt_day(day)} — {fmt_time(clock[:5])}" if clock else fmt_day(day)


def fmt_time(hhmm: str) -> str:
    total = _minutes(hhmm)
    hours, minutes = total // 60, total % 60
    suffix = "ص" if hours < 12 else "م"
    return f"{(hours - 1) % 12 + 1}:{minutes:02d} {suffix}"


def csrf_token() -> str:
    if "csrf" not in session:
        session["csrf"] = secrets.token_urlsafe(24)
    return session["csrf"]


app.jinja_env.globals.update(
    csrf_token=csrf_token, fmt_day=fmt_day, fmt_time=fmt_time, fmt_datetime=fmt_datetime,
    weekday_names=WEEKDAY_NAMES,
)


@app.before_request
def check_csrf():
    if request.method == "POST":
        sent = request.form.get("csrf_token", "")
        if not sent or not hmac.compare_digest(sent, session.get("csrf", "")):
            abort(400, "انتهت صلاحية الصفحة. برجاء إعادة تحميلها والمحاولة مرة أخرى.")


@app.after_request
def security_headers(response):
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "DENY")
    response.headers.setdefault("Referrer-Policy", "same-origin")
    if request.path.startswith(("/admin", "/ticket")):
        response.headers["Cache-Control"] = "no-store"
    return response


# ============================================================
#  تنظيف البيانات والتحقق
# ============================================================
_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
_DIACRITICS = re.compile(r"[ً-ْـ]")
_ARABIC_NAME = re.compile(r"^[ء-ي ]+$")
_SCHOOL = re.compile(r"^[ء-ي0-9A-Za-z .()\-/]+$")
_PHONE = re.compile(r"^01[0125]\d{8}$")


def clean_text(value: str | None) -> str:
    value = _DIACRITICS.sub("", (value or "").translate(_DIGITS))
    return " ".join(value.split())


def clean_digits(value: str | None) -> str:
    return re.sub(r"\D", "", (value or "").translate(_DIGITS))


def parse_national_id(nid: str):
    """يعيد (تاريخ الميلاد، النوع) أو None إذا كان الرقم غير صحيح."""
    if not re.fullmatch(r"[23]\d{13}", nid):
        return None
    year = (1900 if nid[0] == "2" else 2000) + int(nid[1:3])
    try:
        birth = date(year, int(nid[3:5]), int(nid[5:7]))
    except ValueError:
        return None
    if birth > today_cairo():
        return None
    return birth, ("ذكر" if int(nid[12]) % 2 else "أنثى")


def validate_registration(form) -> tuple[dict, dict]:
    data, errors = {}, {}

    data["student_name"] = clean_text(form.get("student_name"))
    if not data["student_name"]:
        errors["student_name"] = "اسم الطالب مطلوب."
    elif not _ARABIC_NAME.match(data["student_name"]):
        errors["student_name"] = "اكتب الاسم بالحروف العربية فقط."
    elif len(data["student_name"].split()) < 4:
        errors["student_name"] = "برجاء كتابة الاسم رباعيًا."
    elif len(data["student_name"]) > 120:
        errors["student_name"] = "الاسم أطول من المسموح."

    data["student_nid"] = clean_digits(form.get("student_nid"))
    parsed = parse_national_id(data["student_nid"])
    if not data["student_nid"]:
        errors["student_nid"] = "الرقم القومي للطالب مطلوب."
    elif parsed is None or data["student_nid"][0] != "3":
        errors["student_nid"] = "الرقم القومي للطالب غير صحيح (14 رقمًا من شهادة الميلاد)."
    else:
        data["birth_date"], data["gender"] = parsed[0].isoformat(), parsed[1]

    for field, choices, label in (
        ("administration", ADMINISTRATIONS, "الإدارة التعليمية"),
        ("grade", GRADE_STAGE, "الصف الدراسي"),
        ("disability", DISABILITY_RULES, "نوع الإعاقة"),
        ("relation", RELATIONS, "صلة القرابة"),
    ):
        data[field] = (form.get(field) or "").strip()
        if data[field] not in choices:
            errors[field] = f"برجاء اختيار {label}."
    data["stage"] = GRADE_STAGE.get(data["grade"], "")

    data["school"] = clean_text(form.get("school"))
    if not data["school"]:
        errors["school"] = "اسم المدرسة مطلوب."
    elif not _SCHOOL.match(data["school"]) or len(data["school"]) > 150:
        errors["school"] = "اسم المدرسة يحتوي على رموز غير مسموحة."

    data["guardian_name"] = clean_text(form.get("guardian_name"))
    if not data["guardian_name"]:
        errors["guardian_name"] = "اسم ولي الأمر مطلوب."
    elif not _ARABIC_NAME.match(data["guardian_name"]):
        errors["guardian_name"] = "اكتب الاسم بالحروف العربية فقط."
    elif len(data["guardian_name"].split()) < 3 or len(data["guardian_name"]) > 120:
        errors["guardian_name"] = "برجاء كتابة اسم ولي الأمر ثلاثيًا على الأقل."

    data["phone"] = clean_digits(form.get("phone"))
    if not _PHONE.match(data["phone"]):
        errors["phone"] = "رقم الهاتف غير صحيح. الصيغة: 01XXXXXXXXX"

    data["whatsapp"] = clean_digits(form.get("whatsapp"))
    if data["whatsapp"] and not _PHONE.match(data["whatsapp"]):
        errors["whatsapp"] = "رقم واتساب غير صحيح. الصيغة: 01XXXXXXXXX"

    for field in ("guardian_nid", "father_nid", "mother_nid"):
        data[field] = clean_digits(form.get(field))
    for field in ("father_name", "mother_name"):
        data[field] = clean_text(form.get(field))

    # ولي الأمر هو الأب أو الأم: لا داعي لكتابة البيانات مرتين
    parent = {"الأب": "father", "الأم": "mother"}.get(data["relation"])
    if parent:
        data[f"{parent}_name"] = data[f"{parent}_name"] or data["guardian_name"]
        if not data[f"{parent}_nid"]:
            data[f"{parent}_nid"] = data["guardian_nid"]
        elif not data["guardian_nid"]:
            data["guardian_nid"] = data[f"{parent}_nid"]

    for field, label in (
        ("guardian_nid", "ولي الأمر"), ("father_nid", "الأب"), ("mother_nid", "الأم"),
    ):
        if not data[field]:
            errors[field] = f"الرقم القومي لـ{label} مطلوب."
        elif parse_national_id(data[field]) is None:
            errors[field] = f"الرقم القومي لـ{label} غير صحيح (14 رقمًا)."
        elif data[field] == data["student_nid"]:
            errors[field] = f"الرقم القومي لـ{label} مطابق لرقم الطالب."
    if data["father_nid"] and data["father_nid"] == data["mother_nid"]:
        errors.setdefault("mother_nid", "الرقم القومي للأم مطابق لرقم الأب.")

    for field, label in (("father_name", "الأب"), ("mother_name", "الأم")):
        if not data[field]:
            errors[field] = f"اسم {label} مطلوب."
        elif not _ARABIC_NAME.match(data[field]) or len(data[field]) > 120:
            errors[field] = f"اكتب اسم {label} بالحروف العربية فقط."
        elif len(data[field].split()) < 3:
            errors[field] = f"برجاء كتابة اسم {label} ثلاثيًا على الأقل."

    # بيانات الحالة
    data["iq_score"] = None
    raw_iq = clean_digits(form.get("iq_score"))
    if len(raw_iq) > 3:
        errors["iq_score"] = "درجة الذكاء يجب أن تكون بين 0 و200."
    elif "disability" not in errors:
        data["iq_score"], iq_error = check_iq(data["disability"], raw_iq)
        if iq_error:
            errors["iq_score"] = iq_error

    for field, limit, label, required in (
        ("condition_description", 1000, "وصف الحالة", False),
        ("report_authority", 200, "جهة إصدار التقرير", True),
        ("city", 100, "المدينة أو المركز", True),
        ("address", 300, "العنوان", True),
    ):
        data[field] = clean_text(form.get(field))
        if required and not data[field]:
            errors[field] = f"{label}: هذا الحقل مطلوب."
        elif len(data[field]) > limit:
            errors[field] = f"{label} أطول من المسموح."

    data["report_date"] = (form.get("report_date") or "").strip()
    if not data["report_date"]:
        errors["report_date"] = "تاريخ التقرير مطلوب."
    else:
        try:
            if date.fromisoformat(data["report_date"]) > today_cairo():
                errors["report_date"] = "تاريخ التقرير لا يمكن أن يكون في المستقبل."
        except ValueError:
            errors["report_date"] = "تاريخ التقرير غير صحيح."

    data["email"] = (form.get("email") or "").strip().lower()
    if data["email"] and (
        len(data["email"]) > 150 or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", data["email"])
    ):
        errors["email"] = "البريد الإلكتروني غير صحيح."

    return data, errors


# ============================================================
#  الصفحات العامة
# ============================================================
def _seconds_until(value: str) -> int:
    if not value:
        return 0
    target = datetime.strptime(value, "%Y-%m-%dT%H:%M")
    return max(0, int((target - now_cairo()).total_seconds()))


def render_form(form, errors, status=200):
    settings = get_settings()
    open_day = find_open_day(get_db(), settings)
    return (
        render_template(
            "form.html",
            form=form,
            errors=errors,
            open_day=open_day[0] if open_day else None,
            # سبب الإغلاق: إيقاف من المسؤول أو اكتمال كل المواعيد
            closed_reason=None if open_day else ("admin" if not settings["is_open"] else "full"),
            closed_message=settings["closed_message"] or DEFAULT_CLOSED_MESSAGE,
            open_at=settings["open_at"] if not settings["is_open"] else "",
            open_in_seconds=_seconds_until(settings["open_at"]) if not settings["is_open"] else 0,
            announcement=settings["announcement"],
            administrations=ADMINISTRATIONS,
            disabilities={
                name: {"iq": iq_range_label(name), "req": rule[2]}
                for name, rule in DISABILITY_RULES.items()
            },
            notes=IMPORTANT_NOTES,
            stages=STAGES,
            relations=RELATIONS,
        ),
        status,
    )


@app.get("/")
def index():
    return render_form({}, {})


@app.post("/")
def register():
    if request.form.get("website"):  # حقل مخفي لصد الروبوتات
        return redirect(url_for("index"))

    data, errors = validate_registration(request.form)
    if errors:
        return render_form(request.form, errors, 422)

    conn = get_db()
    settings = get_settings()
    conn.execute("BEGIN IMMEDIATE")
    try:
        # الطلب الموقوف أو القديم (انتهى موعده) أو المحذوف لا يمنع التسجيل من جديد
        existing = conn.execute(
            "SELECT token, phone FROM registrations"
            " WHERE student_nid = ? AND status = 'active' AND day >= ? ORDER BY id DESC",
            (data["student_nid"], today_cairo().isoformat()),
        ).fetchone()
        if existing is not None:
            conn.execute("ROLLBACK")
            if existing["phone"] == data["phone"]:
                flash("هذا الطالب مسجل من قبل، وهذا موعده.")
                return redirect(url_for("ticket", token=existing["token"]))
            errors["student_nid"] = "هذا الطالب له طلب سارٍ بالفعل. يمكنك معرفة الموعد من صفحة الاستعلام."
            return render_form(request.form, errors, 409)

        open_day = find_open_day(conn, settings)
        if open_day is None:
            conn.execute("ROLLBACK")
            return render_form(request.form, {}, 409)

        day, capacity = open_day
        used = {
            r["seq"]
            for r in conn.execute("SELECT seq FROM registrations WHERE day = ?", (day.isoformat(),))
        }
        seq = next(n for n in range(1, len(used) + 2) if n not in used)
        slot_start, slot_end = slot_for(seq, capacity, settings)
        token = secrets.token_urlsafe(16)
        record = {key: (value if value != "" else None) for key, value in data.items()}
        record.update(
            token=token, code=f"{day:%m%d}-{seq:03d}", day=day.isoformat(), seq=seq,
            slot_start=slot_start, slot_end=slot_end, created_at=now_cairo_str(),
        )
        columns = ", ".join(record)
        conn.execute(
            f"INSERT INTO registrations ({columns}) VALUES ({', '.join('?' * len(record))})",
            list(record.values()),
        )
        conn.execute("COMMIT")
    except INTEGRITY_ERRORS:
        if conn.in_transaction:
            conn.execute("ROLLBACK")
        errors["student_nid"] = "تعذر إتمام التسجيل. برجاء المحاولة مرة أخرى."
        return render_form(request.form, errors, 409)
    except Exception:
        if conn.in_transaction:
            conn.execute("ROLLBACK")
        raise
    return redirect(url_for("ticket", token=token))


def barcode_svg(code: str) -> Markup:
    """باركود Code 128 برقم التسجيل (SVG مضمّن في الصفحة)."""
    buffer = io.BytesIO()
    Code128(code, writer=SVGWriter()).write(
        buffer,
        {"write_text": False, "module_width": 0.42, "module_height": 16, "quiet_zone": 3},
    )
    svg = buffer.getvalue().decode("utf-8")
    return Markup(svg[svg.index("<svg"):])


@app.get("/ticket/<token>")
def ticket(token: str):
    row = get_db().execute("SELECT * FROM registrations WHERE token = ?", (token,)).fetchone()
    if row is None:
        abort(404)
    rule = DISABILITY_RULES.get(row["disability"], (None, None, "", []))
    return render_template(
        "ticket.html",
        r=row,
        documents=GENERAL_DOCUMENTS + rule[3],
        requirements=rule[2],
        barcode=barcode_svg(row["code"]),
        year=ACADEMIC_YEAR_LABEL,
        stopped=row["status"] == STATUS_STOPPED,
        expired=row["day"] < today_cairo().isoformat(),
    )


@app.route("/lookup", methods=["GET", "POST"])
def lookup():
    error = None
    if request.method == "POST":
        row = get_db().execute(
            "SELECT token FROM registrations WHERE student_nid = ? AND phone = ?"
            " ORDER BY (status = 'active') DESC, id DESC",
            (clean_digits(request.form.get("student_nid")), clean_digits(request.form.get("phone"))),
        ).fetchone()
        if row is not None:
            return redirect(url_for("ticket", token=row["token"]))
        error = "لا يوجد تسجيل بهذه البيانات. تأكد من الرقم القومي للطالب ورقم الهاتف المسجل."
    return render_template("lookup.html", error=error, form=request.form)


# ============================================================
#  لوحة الإدارة
# ============================================================
def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("admin"):
            return redirect(url_for("admin_login"))
        return view(*args, **kwargs)

    return wrapped


@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    password = os.environ.get("ADMIN_PASSWORD", "")
    error = None
    if not password:
        error = "لم يتم ضبط كلمة مرور الإدارة. أضف ADMIN_PASSWORD في ملف .env على الخادم."
    elif request.method == "POST":
        if hmac.compare_digest(request.form.get("password", "").encode(), password.encode()):
            session.clear()
            session["admin"] = True
            session.permanent = True
            return redirect(url_for("admin_home"))
        _time.sleep(1.5)
        error = "كلمة المرور غير صحيحة."
    return render_template("admin_login.html", error=error)


@app.post("/admin/logout")
def admin_logout():
    session.clear()
    return redirect(url_for("admin_login"))


@app.get("/admin")
@admin_required
def admin_home():
    conn = get_db()
    settings = get_settings()
    today = today_cairo()
    counts = {
        r["day"]: r["n"]
        for r in conn.execute(
            "SELECT day, COUNT(*) AS n FROM registrations WHERE status = 'active' GROUP BY day"
        )
    }
    overrides = {r["day"]: r for r in conn.execute("SELECT * FROM day_overrides")}

    days = set(counts)
    offset = 0
    upcoming = 0
    while upcoming < 9 and offset <= LOOKAHEAD_DAYS:
        day = today + timedelta(days=offset)
        if day.weekday() in settings["weekdays"]:
            days.add(day.isoformat())
            upcoming += 1
        offset += 1

    rows = []
    for iso in sorted(days):
        override = overrides.get(iso)
        rows.append({
            "day": iso,
            "past": iso < today.isoformat(),
            "count": counts.get(iso, 0),
            "capacity": day_capacity(conn, date.fromisoformat(iso), settings),
            "custom": override["capacity"] if override is not None else None,
            "closed": bool(override["closed"]) if override is not None else False,
        })

    open_day = find_open_day(conn, settings)
    stats = conn.execute(
        """SELECT
               COALESCE(SUM(CASE WHEN status = 'active' AND day >= ? THEN 1 ELSE 0 END), 0) AS upcoming,
               COALESCE(SUM(CASE WHEN status = 'active' AND day = ? THEN 1 ELSE 0 END), 0) AS today,
               COALESCE(SUM(CASE WHEN received_at LIKE ? THEN 1 ELSE 0 END), 0) AS received_today,
               COALESCE(SUM(CASE WHEN received_at IS NOT NULL THEN 1 ELSE 0 END), 0) AS received
           FROM registrations""",
        (today.isoformat(), today.isoformat(), today.isoformat() + "%"),
    ).fetchone()
    return render_template(
        "admin.html",
        settings=settings,
        stats=stats,
        rows=rows,
        total=sum(counts.values()),
        open_day=open_day[0] if open_day else None,
        weekday_order=WEEKDAY_ORDER,
    )


@app.post("/admin/settings")
@admin_required
def admin_settings():
    form = request.form
    try:
        weekdays = sorted({int(x) for x in form.getlist("weekdays") if 0 <= int(x) <= 6})
        capacity = int(form.get("capacity", ""))
        slot_minutes = int(form.get("slot_minutes", ""))
        start, end = form.get("start", ""), form.get("end", "")
        if not (re.fullmatch(r"\d{2}:\d{2}", start) and re.fullmatch(r"\d{2}:\d{2}", end)):
            raise ValueError
        if not (1 <= capacity <= 5000 and 5 <= slot_minutes <= 120 and _minutes(start) < _minutes(end)):
            raise ValueError
    except ValueError:
        flash("بيانات الإعدادات غير صحيحة. راجع العدد والمواعيد.")
        return redirect(url_for("admin_home"))

    values = {
        "weekdays": ",".join(map(str, weekdays)),
        "capacity": str(capacity),
        "start": start,
        "end": end,
        "slot_minutes": str(slot_minutes),
    }
    save_settings(values)
    flash("تم حفظ الإعدادات.")
    return redirect(url_for("admin_home"))


@app.post("/admin/intake")
@admin_required
def admin_intake():
    """فتح/إيقاف استقبال الطلبات ورسائل أولياء الأمور."""
    form = request.form
    action = form.get("action", "save")
    open_at = (form.get("open_at") or "").strip()[:16]
    if open_at:
        try:
            datetime.strptime(open_at, "%Y-%m-%dT%H:%M")
        except ValueError:
            flash("موعد الفتح التلقائي غير صحيح.")
            return redirect(url_for("admin_home"))
    values = {
        "closed_message": (form.get("closed_message") or "").strip()[:500],
        "announcement": (form.get("announcement") or "").strip()[:500],
        "open_at": open_at,
    }
    if action == "close":
        values["is_open"] = "0"
        flash("تم إيقاف استقبال الطلبات. يظهر الآن لأولياء الأمور أن التسجيل مغلق.")
    elif action == "open":
        values.update(is_open="1", open_at="")
        flash("تم فتح استقبال الطلبات.")
    else:
        flash("تم حفظ الرسائل.")
    still_open = values["is_open"] == "1" if "is_open" in values else get_settings()["is_open"]
    if still_open:
        values["open_at"] = ""  # الفتح التلقائي لا معنى له والتسجيل مفتوح
    save_settings(values)
    return redirect(url_for("admin_home"))


def _parse_day(value: str) -> str:
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError:
        abort(404)


@app.post("/admin/day/<day>")
@admin_required
def admin_day_update(day: str):
    day = _parse_day(day)
    raw = clean_digits(request.form.get("capacity"))
    capacity = min(int(raw), 5000) if raw else None
    closed = 1 if request.form.get("closed") else 0
    get_db().execute(
        "INSERT INTO day_overrides (day, capacity, closed) VALUES (?, ?, ?)"
        " ON CONFLICT (day) DO UPDATE SET capacity = excluded.capacity, closed = excluded.closed",
        (day, capacity, closed),
    )
    flash(f"تم تحديث يوم {fmt_day(day)}.")
    return redirect(url_for("admin_home"))


def _query_registrations(day: str | None, q: str = ""):
    sql, params = "SELECT * FROM registrations WHERE 1=1", []
    if day:
        sql += " AND day = ?"
        params.append(day)
    if q:
        like = f"%{q}%"
        sql += " AND (student_name LIKE ? OR student_nid LIKE ? OR phone LIKE ? OR code LIKE ?)"
        params += [like, like, like, like]
    return get_db().execute(sql + " ORDER BY day, seq", params).fetchall()


@app.get("/admin/list")
@admin_required
def admin_list():
    day = _parse_day(request.args["day"]) if request.args.get("day") else None
    q = clean_text(request.args.get("q"))
    rows = _query_registrations(day, q)
    return render_template(
        "admin_list.html",
        rows=rows,
        day=day,
        q=q,
        today=today_cairo().isoformat(),
        # مسح باركود: نتيجة واحدة مطابقة لرقم التسجيل تُعرض كبطاقة استلام سريعة
        scanned=next((r for r in rows if r["code"] == q), None) if q else None,
    )


@app.post("/admin/delete/<int:reg_id>")
@admin_required
def admin_delete(reg_id: int):
    get_db().execute("DELETE FROM registrations WHERE id = ?", (reg_id,))
    flash("تم حذف التسجيل.")
    return _back_to_list()


def _back_to_list():
    target = request.form.get("next", "")
    if not target.startswith("/admin/list"):
        target = url_for("admin_list")
    return redirect(target)


@app.post("/admin/receive/<int:reg_id>")
@admin_required
def admin_receive(reg_id: int):
    """تسجيل استلام أوراق الطالب (أو إلغاؤه)."""
    conn = get_db()
    row = conn.execute("SELECT received_at FROM registrations WHERE id = ?", (reg_id,)).fetchone()
    if row is not None:
        received = None if row["received_at"] else now_cairo_str()
        conn.execute("UPDATE registrations SET received_at = ? WHERE id = ?", (received, reg_id))
        flash("تم تسجيل استلام الأوراق." if received else "تم إلغاء تسجيل الاستلام.")
    return _back_to_list()


@app.post("/admin/stop/<int:reg_id>")
@admin_required
def admin_stop(reg_id: int):
    get_db().execute(
        "UPDATE registrations SET status = 'stopped', stopped_at = ? WHERE id = ?",
        (now_cairo_str(), reg_id),
    )
    flash("تم إيقاف الطلب. يمكن للطالب التسجيل من جديد.")
    return _back_to_list()


@app.post("/admin/resume/<int:reg_id>")
@admin_required
def admin_resume(reg_id: int):
    conn = get_db()
    conn.execute("BEGIN IMMEDIATE")
    row = conn.execute("SELECT student_nid FROM registrations WHERE id = ?", (reg_id,)).fetchone()
    other = row and conn.execute(
        "SELECT 1 FROM registrations"
        " WHERE student_nid = ? AND status = 'active' AND day >= ? AND id != ?",
        (row["student_nid"], today_cairo().isoformat(), reg_id),
    ).fetchone()
    if row and not other:
        conn.execute(
            "UPDATE registrations SET status = 'active', stopped_at = NULL WHERE id = ?", (reg_id,)
        )
    conn.execute("COMMIT")
    flash("لا يمكن إعادة التفعيل: للطالب طلب سارٍ آخر." if other else "تمت إعادة تفعيل الطلب.")
    return _back_to_list()


EXPORT_COLUMNS = [
    ("م", 6, lambda r, i: i),
    ("رقم التسجيل", 14, lambda r, i: r["code"]),
    ("الحالة", 10, lambda r, i: "موقوف" if r["status"] == STATUS_STOPPED else "ساري"),
    ("استلام الأوراق", 18, lambda r, i: r["received_at"] or "لم يتم"),
    ("يوم الحضور", 26, lambda r, i: fmt_day(r["day"])),
    ("التاريخ", 13, lambda r, i: r["day"]),
    ("رقم الدور", 10, lambda r, i: r["seq"]),
    ("من", 11, lambda r, i: fmt_time(r["slot_start"])),
    ("إلى", 11, lambda r, i: fmt_time(r["slot_end"])),
    ("اسم الطالب", 32, lambda r, i: r["student_name"]),
    ("الرقم القومي للطالب", 20, lambda r, i: r["student_nid"]),
    ("تاريخ الميلاد", 14, lambda r, i: r["birth_date"]),
    ("النوع", 8, lambda r, i: r["gender"]),
    ("الإدارة التعليمية", 20, lambda r, i: r["administration"]),
    ("المدرسة", 30, lambda r, i: r["school"]),
    ("المرحلة", 14, lambda r, i: r["stage"]),
    ("الصف", 22, lambda r, i: r["grade"]),
    ("نوع الإعاقة", 26, lambda r, i: r["disability"]),
    ("درجة الذكاء", 12, lambda r, i: r["iq_score"] if r["iq_score"] is not None else ""),
    ("وصف الحالة", 36, lambda r, i: r["condition_description"] or ""),
    ("تاريخ التقرير", 14, lambda r, i: r["report_date"] or ""),
    ("جهة إصدار التقرير", 26, lambda r, i: r["report_authority"] or ""),
    ("اسم ولي الأمر", 30, lambda r, i: r["guardian_name"]),
    ("صلة القرابة", 16, lambda r, i: r["relation"]),
    ("رقم الهاتف", 15, lambda r, i: r["phone"]),
    ("واتساب", 15, lambda r, i: r["whatsapp"] or ""),
    ("البريد الإلكتروني", 26, lambda r, i: r["email"] or ""),
    ("الرقم القومي لولي الأمر", 20, lambda r, i: r["guardian_nid"] or ""),
    ("اسم الأب", 30, lambda r, i: r["father_name"] or ""),
    ("الرقم القومي للأب", 20, lambda r, i: r["father_nid"] or ""),
    ("اسم الأم", 30, lambda r, i: r["mother_name"] or ""),
    ("الرقم القومي للأم", 20, lambda r, i: r["mother_nid"] or ""),
    ("المدينة أو المركز", 18, lambda r, i: r["city"] or ""),
    ("العنوان", 34, lambda r, i: r["address"] or ""),
    ("وقت التسجيل", 18, lambda r, i: r["created_at"]),
]


def _safe_cell(value):
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _write_sheet(workbook: Workbook, title: str, rows) -> None:
    sheet = workbook.create_sheet(title[:31])
    sheet.sheet_view.rightToLeft = True
    sheet.append([column[0] for column in EXPORT_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0D2C54")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, row in enumerate(rows, start=1):
        sheet.append([_safe_cell(column[2](row, index)) for column in EXPORT_COLUMNS])
    for index, column in enumerate(EXPORT_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column[1]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


@app.get("/admin/export")
@admin_required
def admin_export():
    day = _parse_day(request.args["day"]) if request.args.get("day") else None
    rows = _query_registrations(day)
    workbook = Workbook()
    workbook.remove(workbook.active)
    if day:
        _write_sheet(workbook, day, rows)
        filename = f"inclusion-requests-{day}.xlsx"
    else:
        _write_sheet(workbook, "كل الأيام", rows)
        for iso in sorted({r["day"] for r in rows}):
            _write_sheet(workbook, iso, [r for r in rows if r["day"] == iso])
        filename = "inclusion-requests-all.xlsx"
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


ERROR_MESSAGES = {
    404: "الصفحة غير موجودة.",
    405: "طريقة الطلب غير مسموحة.",
    413: "حجم البيانات المرسلة أكبر من المسموح.",
    500: "حدث خطأ غير متوقع. برجاء المحاولة مرة أخرى بعد قليل.",
}


@app.get("/healthz")
def healthz():
    get_db().execute("SELECT 1")
    return "ok"


@app.errorhandler(400)
@app.errorhandler(404)
@app.errorhandler(405)
@app.errorhandler(413)
@app.errorhandler(500)
def handle_error(error):
    code = getattr(error, "code", 500)
    message = ERROR_MESSAGES.get(code) or getattr(error, "description", ERROR_MESSAGES[500])
    return render_template("message.html", message=message), code


init_db()

if USE_POSTGRES and not os.environ.get("SECRET_KEY"):
    _conn = connect()
    _conn.execute(
        "INSERT INTO settings (key, value) VALUES ('secret_key', ?) ON CONFLICT (key) DO NOTHING",
        (secrets.token_hex(32),),
    )
    app.config["SECRET_KEY"] = _conn.execute(
        "SELECT value FROM settings WHERE key = 'secret_key'"
    ).fetchone()[0]
    _conn.close()

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", 5050)), debug=False)
